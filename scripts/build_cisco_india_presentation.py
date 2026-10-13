#!/usr/bin/env python3
"""Build Cisco India Report as an Altera Strategic Account Plan–style PPTX."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# Altera Light template palette (approx. Altera_Light_Clean / Sep2026 Aptos)
NAVY = RGBColor(0x0B, 0x2B, 0x45)
ALTERA_BLUE = RGBColor(0x00, 0x7A, 0xCC)
ACCENT = RGBColor(0x00, 0xA3, 0xE0)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1A, 0x1A, 0x1A)
GRAY = RGBColor(0x5C, 0x6B, 0x7A)
MUTED = RGBColor(0x8A, 0x96, 0xA3)
SWOT_S = RGBColor(0x1F, 0x7A, 0x4C)
SWOT_W = RGBColor(0xB0, 0x35, 0x3A)
SWOT_O = RGBColor(0x00, 0x7A, 0xCC)
SWOT_T = RGBColor(0xC4, 0x7A, 0x12)
ROW_ALT = RGBColor(0xEE, 0xF3, 0xF8)
HEADER_BG = RGBColor(0x0B, 0x2B, 0x45)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Liberation Sans"


def set_run_font(run, size=14, bold=False, color=DARK, font_name=FONT):
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    # Force latin/ea font for broader PPT compatibility
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "ea", "cs"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = rPr.makeelement(qn(f"a:{tag}"), {})
            rPr.append(el)
        el.set("typeface", font_name)


def add_rect(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, *, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}[anchor])
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, *, size=13, color=DARK, bullet=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        p.level = 0
        if bullet:
            pPr = p._p.get_or_add_pPr()
            buFont = pPr.makeelement(qn("a:buFont"), {"typeface": FONT})
            buChar = pPr.makeelement(qn("a:buChar"), {"char": "•"})
            # clear existing bullet nodes then append
            for child in list(pPr):
                if child.tag.endswith("buNone") or child.tag.endswith("buFont") or child.tag.endswith("buChar") or child.tag.endswith("buAutoNum"):
                    pPr.remove(child)
            pPr.insert(0, buChar)
            pPr.insert(0, buFont)
        run = p.add_run()
        run.text = item
        set_run_font(run, size=size, color=color)
    return box


def footer(slide, page: int | None = None, total: int | None = None):
    add_rect(slide, Inches(0), Inches(7.15), SLIDE_W, Inches(0.35), NAVY)
    add_text_box(
        slide,
        Inches(0.4),
        Inches(7.18),
        Inches(6),
        Inches(0.28),
        "Altera Confidential",
        size=10,
        color=WHITE,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    if page is not None and total is not None:
        add_text_box(
            slide,
            Inches(10.5),
            Inches(7.18),
            Inches(2.4),
            Inches(0.28),
            f"{page} / {total}",
            size=10,
            color=WHITE,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
        )


def header_bar(slide, section_label: str):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), ACCENT)
    add_text_box(
        slide,
        Inches(0.5),
        Inches(0.18),
        Inches(8),
        Inches(0.3),
        section_label,
        size=11,
        bold=True,
        color=ALTERA_BLUE,
    )


def title_block(slide, title: str, subtitle: str | None = None):
    add_text_box(slide, Inches(0.5), Inches(0.45), Inches(12.2), Inches(0.55), title, size=28, bold=True, color=NAVY)
    if subtitle:
        add_text_box(slide, Inches(0.5), Inches(0.95), Inches(12.2), Inches(0.35), subtitle, size=13, color=GRAY)


def add_table(slide, left, top, width, height, headers, rows, col_widths=None):
    table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), left, top, width, height)
    table = table_shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = w
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = ""
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = h
        set_run_font(run, size=10, bold=True, color=WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = HEADER_BG
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(val)
            set_run_font(run, size=9, color=DARK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if i % 2 == 0 else ROW_ALT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    return table_shape


def blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # blank


def build(out_path: Path):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    slides_meta: list[str] = []

    # 1. Title
    s = blank_slide(prs)
    add_rect(s, Inches(0), Inches(0), Inches(0.18), SLIDE_H, ACCENT)
    add_rect(s, Inches(0), Inches(0), SLIDE_W, Inches(0.08), ALTERA_BLUE)
    add_text_box(s, Inches(0.7), Inches(1.6), Inches(11.5), Inches(0.4), "STRATEGIC ACCOUNT PLAN", size=14, bold=True, color=ALTERA_BLUE)
    add_text_box(s, Inches(0.7), Inches(2.1), Inches(11.5), Inches(0.8), "CISCO INDIA REPORT", size=40, bold=True, color=NAVY)
    add_text_box(
        s,
        Inches(0.7),
        Inches(3.0),
        Inches(11.5),
        Inches(0.5),
        "AP-INDIA Region · Path to #1 contribution",
        size=18,
        color=GRAY,
    )
    add_text_box(
        s,
        Inches(0.7),
        Inches(4.0),
        Inches(6),
        Inches(1.6),
        "SAM (WW): Jacques Cesbron (Acting)\n"
        "Executive Sponsor: Hua Xue\n"
        "India FSE: Manoj Gopalan\n"
        "India FAE: Rushikesh Purohit · Kasturi Rangan\n"
        "Date: October 2026",
        size=13,
        color=DARK,
    )
    footer(s)
    slides_meta.append("Title")

    # 2. Agenda (template section nav)
    s = blank_slide(prs)
    header_bar(s, "Agenda")
    title_block(s, "Executive Summary", "Cisco India regional deep-dive")
    agenda = [
        ("1", "Opportunity Pipeline", "Strategic Opportunities"),
        ("2", "SWOT Analysis — India Region", "Customer Landscape"),
        ("3", "Critical Needs and Requirements", "Critical Needs"),
        ("4", "Relationship Mapping", "Influence Map"),
    ]
    y = Inches(1.6)
    for num, title, section in agenda:
        add_rect(s, Inches(0.5), y, Inches(12.3), Inches(1.0), LIGHT_BG)
        add_rect(s, Inches(0.5), y, Inches(0.12), Inches(1.0), ACCENT)
        add_text_box(s, Inches(0.9), y + Inches(0.18), Inches(0.6), Inches(0.6), num, size=28, bold=True, color=ALTERA_BLUE, anchor=MSO_ANCHOR.MIDDLE)
        add_text_box(s, Inches(1.7), y + Inches(0.18), Inches(8), Inches(0.4), title, size=20, bold=True, color=NAVY)
        add_text_box(s, Inches(1.7), y + Inches(0.55), Inches(8), Inches(0.3), f"Account Plan section: {section}", size=12, color=GRAY)
        y += Inches(1.15)
    footer(s)
    slides_meta.append("Agenda")

    # 3. Coverage snapshot
    s = blank_slide(prs)
    header_bar(s, "Customer Landscape · Coverage")
    title_block(s, "Cisco India coverage snapshot", "CISCO SYSTEMS (INDIA) PRIVATE LIMITED · Bengaluru design center")
    cards = [
        ("$1.53M", "SFDC open pipe\n(3 registered opps)"),
        ("~$139M", "WW Cisco design pipe\n(Oct 2025)"),
        ("30+", "Cisco FPGA engineers\nin Bengaluru"),
        ("10%", "India FAE Cisco FTE\n(Rushikesh)"),
    ]
    x = Inches(0.5)
    for value, label in cards:
        add_rect(s, x, Inches(1.6), Inches(2.9), Inches(1.8), LIGHT_BG)
        add_rect(s, x, Inches(1.6), Inches(2.9), Inches(0.08), ACCENT)
        add_text_box(s, x + Inches(0.2), Inches(1.9), Inches(2.5), Inches(0.55), value, size=28, bold=True, color=NAVY)
        add_text_box(s, x + Inches(0.2), Inches(2.55), Inches(2.5), Inches(0.7), label, size=12, color=GRAY)
        x += Inches(3.15)
    add_rect(s, Inches(0.5), Inches(3.7), Inches(12.3), Inches(2.8), LIGHT_BG)
    add_text_box(s, Inches(0.75), Inches(3.9), Inches(11.8), Inches(0.35), "Key context", size=14, bold=True, color=ALTERA_BLUE)
    add_bullets(
        s,
        Inches(0.75),
        Inches(4.3),
        Inches(11.8),
        Inches(2.0),
        [
            "India is a named Cisco design location on the WW account plan (optical shared with Italy).",
            "CAB completed in India (Aug 2025) + 4-day training (Sep 2025).",
            "Focus: SP routers / ASR & 8000 control-plane FPGA, optical, platform/BMC attach.",
            "Primary Altera owners: Manoj Gopalan (FSE, ~5% FTE) · Rushikesh Purohit (FAE) · Kasturi Rangan (TM/Lavena tech).",
        ],
        size=14,
    )
    footer(s)
    slides_meta.append("Coverage")

    # 4. Opportunity Pipeline
    s = blank_slide(prs)
    header_bar(s, "Strategic Opportunities · Opportunity Pipeline")
    title_block(s, "Cisco India Opportunity Pipeline", "Registered AP-INDIA SFDC + active India-supported programs")
    headers = ["Opportunity", "Device", "Stage", "Status", "Pipe $", "Owners"]
    rows = [
        ["OAM Engine (Lavena MPA)", "Agilex 7", "Define", "25%", "$600K", "Rushikesh / Kasturi"],
        ["Le Tour PH2", "Agilex 5", "Define", "25%", "$480K", "Rushikesh"],
        ["Agena MPA Control Card", "Agilex 5", "Define", "25%", "$450K", "Narasimhan / Rushikesh"],
        ["ASR1k Traffic Manager", "Agilex 7 M", "Active", "DWIN claim WIP", "WW-linked", "Kasturi / Rushikesh / Shekar"],
        ["Firebolt", "Agilex 5 + MAX 10", "Pursuit", "Open", "TBD", "Kasturi"],
        ["Dragon IO FPGA", "IO FPGA", "Engage", "Open", "TBD", "Kasturi"],
        ["India RFQ (control plane)", "Agilex 5", "Quoted", "At risk", "Pricing-gated", "Manoj / Rushikesh / Sina"],
    ]
    add_table(
        s,
        Inches(0.4),
        Inches(1.45),
        Inches(12.5),
        Inches(4.4),
        headers,
        rows,
        col_widths=[Inches(3.1), Inches(1.7), Inches(1.2), Inches(1.5), Inches(1.4), Inches(3.6)],
    )
    footer(s)
    slides_meta.append("Opportunity Pipeline")

    # 5. Pipeline narrative
    s = blank_slide(prs)
    header_bar(s, "Strategic Opportunities · Opportunity Pipeline")
    title_block(s, "Pipeline notes & near-term conversion", "Focus to grow India contribution to WW Cisco pipe")
    add_rect(s, Inches(0.5), Inches(1.55), Inches(6.0), Inches(4.9), LIGHT_BG)
    add_text_box(s, Inches(0.75), Inches(1.75), Inches(5.5), Inches(0.35), "Registered SFDC total: $1.53M", size=16, bold=True, color=NAVY)
    add_bullets(
        s,
        Inches(0.75),
        Inches(2.25),
        Inches(5.5),
        Inches(3.9),
        [
            "Lavena $600K — AMD timing pressure reopening Agilex 7; PCIe/AXI fit analysis with India eng.",
            "Le Tour $480K + Agena $450K — Agilex 5 control/platform path; Agena proto Mar’26 / FCS May’26.",
            "TM / Sausalito (AGMA032) — P1A power-on Oct 12–16 2026; India FAE owning day-to-day APS.",
            "India RFQ — Agilex 5 ~40% above Lattice (~$20/100KLE); golden-part pricing may be required.",
        ],
        size=13,
    )
    add_rect(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(4.9), LIGHT_BG)
    add_rect(s, Inches(6.8), Inches(1.55), Inches(6.0), Inches(0.08), ACCENT)
    add_text_box(s, Inches(7.05), Inches(1.85), Inches(5.5), Inches(0.35), "Near-term conversion focus", size=16, bold=True, color=NAVY)
    add_bullets(
        s,
        Inches(7.05),
        Inches(2.4),
        Inches(5.5),
        Inches(3.7),
        [
            "Claim / register Traffic Manager (Sausalito) DWIN split with WW (Shekar access).",
            "Keep Lavena Agilex 7 alive; clone/track separate TM oppty as needed.",
            "Unblock India RFQ with Agilex 5 pricing strategy.",
            "Close Firebolt (Agilex 5 + MAX 10) and Dragon IO firmware follow-ups.",
        ],
        size=13,
    )
    footer(s)
    slides_meta.append("Pipeline Focus")

    # 6. SWOT
    s = blank_slide(prs)
    header_bar(s, "Customer Landscape · SWOT")
    title_block(s, "SWOT Analysis — India Region (Altera View)", "Control-plane / SP design activity · Bengaluru")
    quads = [
        ("Strengths", SWOT_S, [
            "30+ FPGA eng in Bengaluru owning SP routers, ASR/8000 control path, optical (w/ Italy)",
            "Friendly exec relationship with Vijay Pramod; CAB + training delivered (2025)",
            "Live traction: ASR1k TM on Agilex 7 M; AI/ML tensor story for control/BMC",
            "PQC / secure-boot roadmap vs Microchip security incumbency",
        ]),
        ("Weaknesses", SWOT_W, [
            "Xilinx Artix incumbency on control paths; SmartFusion2 for security",
            "Existing stacks ported to Xilinx / Microchip — high switching cost",
            "Agilex 5 price gap (~40% vs Lattice); weak below ~$8–$20 vs PolarFire/Lattice",
            "Thin coverage: 5% FSE / 10% FAE Cisco FTE vs rising support load",
        ]),
        ("Opportunities", SWOT_O, [
            "Control-plane sockets on 8000 / ASR from Bangalore; Agilex 5/3 + MAX 10 attach",
            "Lavena returning to Agilex 7 after AMD timing issues; TM DWIN + Berkeley ramp",
            "PQC, SEU, AI/ML telemetry demos already in motion with India eng",
            "Cisco manufacturing investment in India; Marco Croci push for India → Altera",
        ]),
        ("Threats", SWOT_T, [
            "AMD Versal; Lattice mid-range; PolarFire security lock-in on CAT8k-class",
            "Datapath → Silicon One / ASICs — FPGA TAM concentrates on control/glue",
            "RFQ pricing losses; low-end secure ~$8 device request closed (no product)",
            "Geopolitical / indigenous-FPGA pressure; WW inventory diverting focus",
        ]),
    ]
    positions = [
        (Inches(0.4), Inches(1.45)),
        (Inches(6.85), Inches(1.45)),
        (Inches(0.4), Inches(4.25)),
        (Inches(6.85), Inches(4.25)),
    ]
    for (title, color, items), (left, top) in zip(quads, positions):
        add_rect(s, left, top, Inches(6.05), Inches(2.55), LIGHT_BG)
        add_rect(s, left, top, Inches(6.05), Inches(0.38), color)
        add_text_box(s, left + Inches(0.2), top + Inches(0.05), Inches(5.6), Inches(0.3), title, size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
        add_bullets(s, left + Inches(0.2), top + Inches(0.5), Inches(5.65), Inches(1.95), items, size=11)
    footer(s)
    slides_meta.append("SWOT")

    # 7. Critical Needs
    s = blank_slide(prs)
    header_bar(s, "Critical Needs")
    title_block(s, "Critical Needs and Requirements", "India engineering asks aligned to WW Cisco key-needs list")
    headers = ["Need", "Why it matters in India", "Status / ask", "Owner"]
    rows = [
        ["Agilex 5 pricing (130K/200K LE)", "RFQ vs Lattice ~$20/100KLE; ~40% gap", "Golden-part pricing required", "Sabrina / Jacques / Manoj"],
        ["Low-end secure device (~$8)", "PolarFire/Lattice sockets dominate", "No product path — risk accepted", "Deepali (closed)"],
        ["TM 220–240 MHz + MCDMA scale", "ASR1k 110–120 MPPS; India day-to-day", "240 MHz experimental; patches open", "Kasturi / Factory APS"],
        ["PQC / secure-boot / programmed parts", "Displace SmartFusion2; Berkeley 2027", "SEU done; black-key / CM signing", "Security eng + India FAE"],
        ["DDR5 on Agilex 7", "Required beyond M-series TM wins", "Not planned on Agilex 7", "Sam / Product"],
        ["Agilex 5/7 loaners + TCL training", "Bengaluru migration evaluation", "Loaners / remote CRDC in flight", "Kasturi / Rushikesh"],
        ["FAE bandwidth for Cisco India", "TM+Lavena+Firebolt overload", "HC #2 NDC/COMM (H1 FY27)", "Rushikesh / India mgmt"],
    ]
    add_table(
        s,
        Inches(0.35),
        Inches(1.4),
        Inches(12.6),
        Inches(5.2),
        headers,
        rows,
        col_widths=[Inches(3.2), Inches(3.5), Inches(3.2), Inches(2.7)],
    )
    footer(s)
    slides_meta.append("Critical Needs")

    # 8. Relationship Mapping — Cisco
    s = blank_slide(prs)
    header_bar(s, "Influence Map · Relationship Mapping")
    title_block(s, "Cisco India stakeholders", "Bengaluru influence map — consensus-driven with Sr. Director sign-off")
    headers = ["Name", "Role", "Stance", "Engagement notes"]
    rows = [
        ["Vijay Pramod Kappagantula", "FPGA Team Director", "Altera Friendly", "Met Aug 2025 (Ganga); SP routing/access/optical/cable FPGA"],
        ["Alap Shah", "Sr. Director, Central Eng", "Strategic", "24 yrs; platform HW & FPGA for optical, core & edge routing"],
        ["Akhil Guha", "Sr. Director, Hardware", "Strategic", "SP access/routing/aggregation HW from Bangalore"],
        ["Manikandan", "Sr. Architect", "Technical", "FPGA architecture; briefing attendee"],
        ["Ramesh / Jasprit / Rakesh /\nHarsh / Kannan / Vaibhav", "FPGA / TM engineering", "Day-to-day", "Lavena, TM timing, Quartus, MCDMA validation"],
    ]
    add_table(
        s,
        Inches(0.4),
        Inches(1.4),
        Inches(12.5),
        Inches(4.0),
        headers,
        rows,
        col_widths=[Inches(3.0), Inches(2.6), Inches(1.7), Inches(5.2)],
    )
    add_text_box(
        s,
        Inches(0.5),
        Inches(5.6),
        Inches(12.3),
        Inches(1.0),
        "Influence path: Team leads / architects → Vijay Pramod (FPGA Director) → Akhil Guha / Alap Shah (Sr. Directors).\n"
        "Technical proof (timing, MCDMA, pricing) must land at working level before director design commit. Optical dual-tracks with Cisco Italy.",
        size=12,
        color=GRAY,
    )
    footer(s)
    slides_meta.append("Cisco Relationships")

    # 9. Relationship Mapping — Altera / WW
    s = blank_slide(prs)
    header_bar(s, "Influence Map · Relationship Mapping")
    title_block(s, "Altera coverage & WW linkage", "Exec relationship map — India region")
    headers = ["Altera", "Role", "Cisco scope", "Responsibility / next steps"]
    rows = [
        ["Rushikesh Purohit", "FAE / FAE Manager", "AP-INDIA · ~10% FTE", "Account lead, oppty owner, briefing owner"],
        ["Manoj Gopalan", "FSE", "AP-INDIA · ~5% FTE", "$250K DW / $1.15M rev quota slice"],
        ["Kasturi Rangan", "FAE (NDC)", "Cisco India programs", "TM / Lavena / Firebolt technical ownership"],
        ["Ganga Gopal", "Leadership", "India exec engagement", "Relationship owner for Vijay Pramod"],
        ["Jacques Cesbron", "SAM (Acting)", "WW Cisco", "Account strategy & air cover"],
        ["Shekar Chandra", "FSE San Jose", "WW ↔ India bridge", "TM DWIN split / US oppty access"],
        ["Hua Xue", "Exec Sponsor", "WW Cisco", "Executive sponsorship"],
        ["Marco Croci", "VP Eng (Cisco side bridge)", "Italy / SJ", "Push Indian engineering to use Altera"],
    ]
    add_table(
        s,
        Inches(0.35),
        Inches(1.4),
        Inches(12.6),
        Inches(5.2),
        headers,
        rows,
        col_widths=[Inches(2.6), Inches(2.4), Inches(2.8), Inches(4.8)],
    )
    footer(s)
    slides_meta.append("Altera Coverage")

    # 10. Closing / sources
    s = blank_slide(prs)
    header_bar(s, "Appendix")
    title_block(s, "Sources & disclaimer", "Aligned to 2025 Altera Strategic Account Plan Template")
    add_bullets(
        s,
        Inches(0.6),
        Inches(1.6),
        Inches(12),
        Inches(4.5),
        [
            "2025 Altera Strategic Account Plan Template (WW Strategic Accounts)",
            "2025 CISCO ACCOUNT PLAN October 25 update",
            "Briefing Document_Cisco_India (Rushikesh Purohit, Jul 2025)",
            "Pipeline Raw Report (SFDC) · India_weekly_update · KasturiRanganWeeklyUpdate",
            "Altera_Status_to_Cisco — Traffic Manager engineering status (Oct 2026)",
            "India_2026Plan / FAE growth notes",
            "Disclaimer: Pipe dollars/stages from Glean/SFDC extracts — confirm live SFDC before external use. WW $139M pipe is account-wide, not India-only.",
        ],
        size=14,
    )
    footer(s)
    slides_meta.append("Sources")

    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        add_text_box(
            slide,
            Inches(10.5),
            Inches(7.18),
            Inches(2.4),
            Inches(0.28),
            f"{i} / {total}",
            size=10,
            color=WHITE,
            align=PP_ALIGN.RIGHT,
            anchor=MSO_ANCHOR.MIDDLE,
        )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    print(f"Wrote {out_path} ({total} slides): {', '.join(slides_meta)}")


if __name__ == "__main__":
    build(Path("/workspace/presentations/Cisco_India_Report_Altera_Account_Plan.pptx"))
