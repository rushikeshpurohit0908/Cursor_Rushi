#!/usr/bin/env python3
"""Generate Agilex 7 + TI AFE7950 DRA PowerPoint (16:9)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

INK = RGBColor(0x0C, 0x16, 0x20)
MUTED = RGBColor(0x4A, 0x5A, 0x66)
PAPER = RGBColor(0xF7, 0xF9, 0xFB)
PANEL = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD0, 0xD8, 0xE0)
RX = RGBColor(0x0E, 0x6B, 0x7A)
RX_LIGHT = RGBColor(0x18, 0xA0, 0xB0)
TX = RGBColor(0xA0, 0x5C, 0x18)
TX_LIGHT = RGBColor(0xC2, 0x77, 0x2A)
FPGA = RGBColor(0x1A, 0x4A, 0x6E)
RFIC = RGBColor(0x1E, 0x5A, 0x48)
STEEL = RGBColor(0x14, 0x2E, 0x44)
ROUTER = RGBColor(0x1B, 0x3F, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CHIP = RGBColor(0xE6, 0xF3, 0xF5)
CHIP_TX = RGBColor(0xF7, 0xEE, 0xE2)
CHIP_RFIC = RGBColor(0xE4, 0xF2, 0xEC)
DARK = RGBColor(0x06, 0x10, 0x18)

W = Inches(13.333)
H = Inches(7.5)
MX = Inches(0.65)
MY = Inches(0.4)

OUT = Path(__file__).resolve().parent / "Agilex7_TI_AFE7950_DRA.pptx"
PNG = Path(__file__).resolve().parent / "block_diagram.png"


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def stroke_shape(shape, color, width_pt=1.0):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def add_rect(slide, left, top, width, height, fill, stroke=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.adjustments[0] = 0.08
    fill_shape(shape, fill)
    if stroke:
        stroke_shape(shape, stroke, 1.0)
    return shape


def add_box(slide, left, top, width, height, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill_shape(shape, fill)
    return shape


def add_text(slide, left, top, width, height, lines, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, (text, size, bold, color, font) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.alignment = align
        if p.runs:
            set_run(p.runs[0], size, bold, color, font)
        p.space_after = Pt(4)
    return box


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def header(slide, eyebrow, title, subtitle=None, dark=False):
    add_box(slide, 0, 0, W, Inches(0.08), RX_LIGHT if not dark else RX)
    ey_c = RX_LIGHT if not dark else RGBColor(0x5F, 0xD0, 0xE0)
    ti_c = STEEL if not dark else WHITE
    su_c = MUTED if not dark else RGBColor(0x9B, 0xB0, 0xC0)
    add_text(slide, MX, MY, Inches(12), Inches(0.3), [(eyebrow.upper(), 11, True, ey_c, "Calibri")])
    add_text(slide, MX, Inches(0.65), Inches(12), Inches(0.55), [(title, 28, False, ti_c, "Georgia")])
    if subtitle:
        add_text(slide, MX, Inches(1.15), Inches(12), Inches(0.4), [(subtitle, 14, False, su_c, "Calibri")])


def card(slide, left, top, width, height, title, bullets, accent=RX):
    add_rect(slide, left, top, width, height, PANEL, LINE)
    add_box(slide, left, top, Inches(0.08), height, accent)
    lines = [(title, 15, True, INK, "Calibri")]
    for b in bullets:
        lines.append((b, 13, False, MUTED, "Calibri"))
    add_text(slide, left + Inches(0.22), top + Inches(0.16), width - Inches(0.35), height - Inches(0.25), lines)


def device_cell(slide, left, top, width, height, title, sub, fill, accent):
    add_rect(slide, left, top, width, height, fill, accent)
    add_text(
        slide,
        left + Inches(0.06),
        top + Inches(0.1),
        width - Inches(0.12),
        height - Inches(0.12),
        [(title, 12, True, INK, "Calibri"), (sub, 10, False, MUTED, "Calibri")],
        align=PP_ALIGN.CENTER,
    )


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H

    # 1 Title
    s = blank(prs)
    set_bg(s, DARK)
    add_box(s, 0, 0, Inches(0.18), H, RX_LIGHT)
    add_text(s, MX, Inches(1.6), Inches(11), Inches(0.35),
             [("ALTERA AGILEX 7  +  TEXAS INSTRUMENTS AFE7950", 12, True, RX_LIGHT, "Calibri")])
    add_text(s, MX, Inches(2.1), Inches(12), Inches(1.3),
             [("Discrete RFIC + FPGA", 40, False, WHITE, "Georgia"),
              ("DRA Architecture", 40, False, WHITE, "Georgia")])
    add_text(s, MX, Inches(3.8), Inches(11), Inches(1.2),
             [("Ku-band uplink & downlink with JESD204C interconnect", 16, False, RGBColor(0x9B, 0xB0, 0xC0), "Calibri"),
              ("9x Agilex 7  |  18x AFE7950  |  Separate RX / TX clusters", 16, False, RGBColor(0x9B, 0xB0, 0xC0), "Calibri")])
    add_text(s, MX, Inches(6.5), Inches(11), Inches(0.3),
             [("Alternative to Direct-RF SoC: RF sampling in TI RFIC, beamforming/PHY in Agilex 7", 12, False, RGBColor(0x7A, 0x93, 0xA6), "Calibri")])

    # 2 Overview
    s = blank(prs)
    set_bg(s, PAPER)
    header(s, "System Overview", "Split roles: RFIC converts, FPGA computes",
           "Same 32/40 beam DRA map as the Direct-RF option — different silicon partition")
    card(s, MX, Inches(1.8), Inches(3.9), Inches(4.6), "TI AFE7950 RFIC",
         ["4 TX + 4 RX (+2 feedback)",
          "12 GSPS DAC / 3 GSPS ADC",
          "Integrated DUC / DDC / NCO / DSA",
          "RF coverage to 12 GHz",
          "JESD204B/C Subclass-1",
          "Space path: AFE7950-SP"], RFIC)
    card(s, Inches(4.75), Inches(1.8), Inches(3.9), Inches(4.6), "Altera Agilex 7",
         ["Digital beamforming network",
          "UL channelizer / PHY framer",
          "DVB-RCS2 / DVB-S2X / NTN",
          "High-speed SerDes (JESD)",
          "HPS control & telemetry",
          "Cluster phase / SYSREF sync"], FPGA)
    card(s, Inches(8.5), Inches(1.8), Inches(4.15), Inches(4.6), "Slot model",
         ["1x Agilex 7 + 2x AFE7950",
          "8 RF channels per slot",
          "RX: 4 slots (32 beams)",
          "TX: 5 slots (40 beams)",
          "Shared LO / SYSREF tree",
          "Router fabric between sides"], ROUTER)

    # 3 Block diagram
    s = blank(prs)
    set_bg(s, DARK)
    if PNG.exists():
        margin = Inches(0.18)
        avail_w = W - 2 * margin
        avail_h = H - 2 * margin
        aspect = 1680 / 1180
        if avail_w / avail_h > aspect:
            pic_h = avail_h
            pic_w = pic_h * aspect
        else:
            pic_w = avail_w
            pic_h = pic_w / aspect
        left = (W - pic_w) / 2
        top = (H - pic_h) / 2
        s.shapes.add_picture(str(PNG), left, top, width=pic_w, height=pic_h)
    else:
        header(s, "Block Diagram", "Full architecture", dark=True)
        add_text(s, MX, Inches(3), Inches(12), Inches(1),
                 [("block_diagram.png missing — run cairosvg first", 16, False, WHITE, "Calibri")])

    # 4 RX
    s = blank(prs)
    set_bg(s, PAPER)
    header(s, "RX Uplink", "14.0–14.5 GHz · 32 beams · Super-het into AFE",
           "Antenna → LNA → Ku→IF mixer → 8× AFE7950 → JESD204C → 4× Agilex 7")
    for i, label in enumerate([
        ("FPGA 1", "2× AFE · Beams 1–8"),
        ("FPGA 2", "2× AFE · Beams 9–16"),
        ("FPGA 3", "2× AFE · Beams 17–24"),
        ("FPGA 4", "2× AFE · Beams 25–32"),
    ]):
        device_cell(s, MX + i * Inches(3.05), Inches(1.85), Inches(2.9), Inches(1.15),
                    label[0], label[1], CHIP, FPGA)
    card(s, MX, Inches(3.3), Inches(5.9), Inches(3.4), "Why super-het on RX?",
         ["AFE7950 RF input tops out near 12 GHz",
          "Ku uplink (14.0–14.5) mixes to IF (~8 GHz class)",
          "LO from LMX2615-SP / LMX2820",
          "AFE handles ADC + DDC/NCO/DSA",
          "Agilex 7 runs DBFN + channelizer + Low-PHY"], RX)
    card(s, Inches(6.9), Inches(3.3), Inches(5.9), Inches(3.4), "FPGA DSP (per slot)",
         ["Digital beamforming: 4-color reuse",
          "8-way spatial reuse",
          "UL channelizer: 4 × 125 MHz / beam",
          "Handoff: DVB-RCS2 / 3GPP-NTN Low-PHY",
          "Egress to onboard router fabric"], FPGA)

    # 5 TX
    s = blank(prs)
    set_bg(s, PAPER)
    header(s, "TX Downlink", "10.7–12.75 GHz · 40 / 32 beams · RF DAC + optional upconvert",
           "5× Agilex 7 → JESD204C → 10× AFE7950 → HPA → Antenna")
    for i, label in enumerate([
        ("FPGA 1", "Beams 1–8"),
        ("FPGA 2", "Beams 9–16"),
        ("FPGA 3", "Beams 17–24"),
        ("FPGA 4", "Beams 25–32"),
        ("FPGA 5 *", "Beams 33–40"),
    ]):
        device_cell(s, MX + i * Inches(2.45), Inches(1.85), Inches(2.3), Inches(1.1),
                    label[0], label[1], CHIP_TX, FPGA)
    card(s, MX, Inches(3.25), Inches(5.9), Inches(3.45), "RF transmit path",
         ["AFE7950 12 GSPS DAC covers much of Ku DL",
          "Direct RF TX up to ~12 GHz",
          "Optional IF→Ku for 12.0–12.75 GHz edge",
          "TX DSA + external HPA front-end",
          "* DVB-S2X-only: idle beams 33–40"], TX)
    card(s, Inches(6.9), Inches(3.25), Inches(5.9), Inches(3.45), "FPGA DSP (per slot)",
         ["DVB-S2X / 3GPP-NTN Low-PHY framer",
          "DBFN: 4-color · 8–10 way spatial",
          "Baseband / IQ to AFE over JESD204C",
          "AFE integrates DUC + RF DAC",
          "2× AFE7950 = 8 DAC channels / slot"], FPGA)

    # 6 Interconnect & sync
    s = blank(prs)
    set_bg(s, PAPER)
    header(s, "Interconnect & Sync", "JESD204C, SYSREF, and fabric")
    card(s, MX, Inches(1.8), Inches(3.9), Inches(4.6), "JESD204C",
         ["Subclass-1 deterministic latency",
          "AFE7950 ↔ Agilex 7 SerDes",
          "8b/10b or 64b/66b modes",
          "Per-slot multi-lane links",
          "Scales with channel BW"], RFIC)
    card(s, Inches(4.75), Inches(1.8), Inches(3.9), Inches(4.6), "Coherent array",
         ["Shared SYSREF distribution",
          "Common LO / sample clock tree",
          "LMX2615-SP or LMX2820",
          "Multi-device phase alignment",
          "Required for DBFN spatial reuse"], FPGA)
    card(s, Inches(8.5), Inches(1.8), Inches(4.15), Inches(4.6), "Router fabric",
         ["Bridges RX Low-PHY → TX framer",
          "Feeder-link stage attach",
          "58 Gbps-class SerDes links",
          "Keeps RX/TX clusters isolated",
          "Control plane via FPGA HPS"], ROUTER)

    # 7 Comparison
    s = blank(prs)
    set_bg(s, PAPER)
    header(s, "Architecture Comparison", "Direct-RF SoC vs Agilex 7 + TI RFIC")
    rows = [
        ("Aspect", "Direct-RF SoC (AGRW)", "Agilex 7 + AFE7950"),
        ("RF conversion", "On-chip Direct-RF ADC/DAC", "TI AFE7950 RFIC"),
        ("Digital DSP", "Same device as converters", "Agilex 7 FPGA"),
        ("RX @ 14 GHz", "1st Nyquist Direct-RF", "Ku→IF super-het + AFE"),
        ("TX @ Ku DL", "Direct-RF DAC", "RF DAC (+ opt. upconvert)"),
        ("Interconnect", "On-chip / SerDes fabric", "JESD204C + fabric"),
        ("Device count", "9 RF SoCs", "9 FPGA + 18 RFIC"),
        ("Flexibility", "Integrated density", "Multi-vendor / SP grade"),
    ]
    table_shape = s.shapes.add_table(len(rows), 3, MX, Inches(1.7), Inches(12.0), Inches(5.0))
    table = table_shape.table
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(4.8)
    table.columns[2].width = Inches(4.8)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            run = cell.text_frame.paragraphs[0].runs[0]
            if r == 0:
                set_run(run, 12, True, WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = STEEL
            else:
                set_run(run, 12, c == 0, INK if c == 0 else MUTED)
                cell.fill.solid()
                cell.fill.fore_color.rgb = PANEL if r % 2 else RGBColor(0xEE, 0xF3, 0xF6)

    # 8 BOM snapshot
    s = blank(prs)
    set_bg(s, PAPER)
    header(s, "BOM Snapshot", "Conservative separate TX/RX path")
    rows = [
        ("Item", "Qty", "Notes"),
        ("Agilex 7 (RX)", "4", "8 beams each via 2× AFE"),
        ("Agilex 7 (TX)", "5", "FPGA 5 optional in DVB-S2X"),
        ("AFE7950 (RX)", "8", "4 RX ch × 8 = 32"),
        ("AFE7950 (TX)", "10", "4 TX ch × 10 = 40"),
        ("LMX2615-SP / LMX2820", "Cluster", "LO + SYSREF distribution"),
        ("Ku→IF mixers (RX)", "32 paths", "Super-het front-end"),
        ("HPA / LNA FE", "Array", "Per-beam RF front-end"),
    ]
    table_shape = s.shapes.add_table(len(rows), 3, MX, Inches(1.7), Inches(12.0), Inches(5.0))
    table = table_shape.table
    table.columns[0].width = Inches(4.0)
    table.columns[1].width = Inches(2.0)
    table.columns[2].width = Inches(6.0)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            run = cell.text_frame.paragraphs[0].runs[0]
            if r == 0:
                set_run(run, 12, True, WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = STEEL
            else:
                set_run(run, 12, c == 0, INK if c == 0 else MUTED)
                cell.fill.solid()
                cell.fill.fore_color.rgb = PANEL if r % 2 else RGBColor(0xEE, 0xF3, 0xF6)

    # 9 Closing
    s = blank(prs)
    set_bg(s, DARK)
    add_box(s, 0, 0, Inches(0.18), H, TX_LIGHT)
    add_text(s, MX, Inches(2.1), Inches(11), Inches(0.35),
             [("SUMMARY", 12, True, RX_LIGHT, "Calibri")])
    add_text(s, MX, Inches(2.55), Inches(12), Inches(1.3),
             [("Agilex 7 digital fabric", 32, False, WHITE, "Georgia"),
              ("meets TI RF sampling", 32, False, WHITE, "Georgia")])
    add_text(s, MX, Inches(4.3), Inches(11), Inches(1.6), [
        ("32-beam Ku UL (super-het) + 40/32-beam Ku DL", 15, False, RGBColor(0x9B, 0xB0, 0xC0), "Calibri"),
        ("JESD204C slot model: 1 FPGA + 2 AFE7950 = 8 channels", 15, False, RGBColor(0x9B, 0xB0, 0xC0), "Calibri"),
        ("Space-ready path via AFE7950-SP + LMX2615-SP", 15, False, RGBColor(0x9B, 0xB0, 0xC0), "Calibri"),
    ])

    prs.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
